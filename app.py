from flask import Flask, render_template, request, redirect, session, url_for, jsonify, send_from_directory, render_template_string
import os, json, bcrypt, dropbox, re, subprocess, pickle
import pandas as pd
from dotenv import load_dotenv
from werkzeug.utils import secure_filename
from PyPDF2 import PdfReader
import docx
import openpyxl
from datetime import datetime
import google.generativeai as genai
from flask_dance.contrib.google import make_google_blueprint, google
from video_summarizer import video_data, summarize, TEMPLATE, SUMMARY_DIR
import firebase_admin
from firebase_admin import credentials, db as admin_db

# ---------------------------- CONFIGURATION ----------------------------
load_dotenv()

app = Flask(__name__)
app.secret_key = 'secret_key'
os.environ['OAUTHLIB_INSECURE_TRANSPORT'] = '1'

UPLOAD_FOLDER = 'uploads'
SUMMARY_FOLDER = 'summaries'
ALLOWED_EXTENSIONS = {'pdf', 'docx', 'xlsx'}

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['SUMMARY_FOLDER'] = SUMMARY_FOLDER
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(SUMMARY_FOLDER, exist_ok=True)

firebase_config = {
    "type": "service_account",
    "project_id": os.getenv("FB_PROJECT_ID"),
    "private_key_id": os.getenv("FB_PRIVATE_KEY_ID"),
    "private_key": os.getenv("FB_PRIVATE_KEY").replace('\\n', '\n'),
    "client_email": os.getenv("FB_CLIENT_EMAIL"),
    "client_id": os.getenv("FB_CLIENT_ID"),
    "auth_uri": os.getenv("FB_AUTH_URI"),
    "token_uri": os.getenv("FB_TOKEN_URI"),
    "auth_provider_x509_cert_url": os.getenv("FB_AUTH_PROVIDER_X509_CERT_URL"),
    "client_x509_cert_url": os.getenv("FB_CLIENT_X509_CERT_URL")
}

# ---------------------------- FIREBASE SETUP ----------------------------
cred = credentials.Certificate(firebase_config)
firebase_admin.initialize_app(cred, {
    'databaseURL': os.getenv("FIREBASE_DATABASE_URL")
})
db_ref = admin_db.reference("/")

# ---------------------------- GOOGLE AUTH ----------------------------
google_bp = make_google_blueprint(
    client_id=os.getenv("GOOGLE_CLIENT_ID"),
    client_secret=os.getenv("GOOGLE_CLIENT_SECRET"),
    redirect_to="google_auth",
    scope=[
        "https://www.googleapis.com/auth/userinfo.profile",
        "https://www.googleapis.com/auth/userinfo.email",
        "openid"
    ]
)
app.register_blueprint(google_bp, url_prefix="/login")

# ---------------------------- DROPBOX & GEMINI ----------------------------
dbx = dropbox.Dropbox(os.getenv("DROPBOX_ACCESS_TOKEN"))
genai.configure(api_key=os.getenv("GEMINI_API_KEY"))

# ---------------------------- HELPER FUNCTIONS ----------------------------
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def clean_department(dept):
    return re.sub(r'\s+', '_', dept.strip().lower())

def extract_text(filepath):
    ext = filepath.rsplit('.', 1)[1].lower()
    text = ""
    if ext == 'pdf':
        reader = PdfReader(filepath)
        for page in reader.pages:
            content = page.extract_text()
            if content:
                text += content
    elif ext == 'docx':
        doc = docx.Document(filepath)
        for para in doc.paragraphs:
            text += para.text + '\n'
    elif ext == 'xlsx':
        wb = openpyxl.load_workbook(filepath)
        for sheet in wb:
            for row in sheet.iter_rows(values_only=True):
                text += ' '.join(str(cell) for cell in row if cell) + '\n'
    return text.strip()

def summarize_text(text):
    try:
        model = genai.GenerativeModel("models/gemini-2.0-flash")
        response = model.generate_content(f"Summarize the following:\n\n{text}")
        return response.text.strip()
    except Exception as e:
        print("Summarize error:", e)
        return "⚠️ Failed to summarize."

# ---------------------------- LOAD PREDICTION MODEL ----------------------------
with open("trained_model.pkl", "rb") as f:
    model = pickle.load(f)

adjustment_factors = {
    "Foggy": 1.10,
    "Rainy": 1.15,
    "Dusty": 1.05
}

# ---------------------------- ROUTES ----------------------------
@app.route('/')
def index():
    return redirect('/login')

@app.route('/signup', methods=['GET', 'POST'])
def signup():
    if request.method == 'POST':
        email = request.form['email']
        username = request.form['username']
        password = request.form['password']
        department = clean_department(request.form['department'])

        users = db_ref.child("users").get() or {}
        if any(u.get("email") == email for u in users.values()):
            return render_template('signup.html', status="Email already registered.", success=False)

        hashed_pw = bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt())
        db_ref.child("users").push({
            "username": username,
            "email": email,
            "password": hashed_pw.decode('utf-8'),
            "role": "user",
            "department": department
        })
        return redirect('/login?status=Account created successfully&success=true')
    return render_template('signup.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    status = request.args.get("status")
    success = request.args.get("success") == "true"

    if request.method == 'POST':
        email = request.form['email']
        password = request.form['password']
        users = db_ref.child("users").get() or {}

        for key, user in users.items():
            if user['email'] == email and bcrypt.checkpw(password.encode('utf-8'), user['password'].encode('utf-8')):
                session.update({
                    'user': email,
                    'username': user.get('username', 'User'),
                    'role': user.get('role', 'user'),
                    'department': user.get('department', ''),
                    'user_id': key
                })
                return redirect('/admin' if user['role'] == 'admin' else '/dashboard')
        return render_template('login.html', status="Invalid credentials", success=False)
    return render_template('login.html', status=status, success=success)

@app.route('/logout')
def logout():
    session.clear()
    return redirect('/login')

@app.route('/google_auth')
def google_auth():
    if not google.authorized:
        return redirect(url_for("google.login"))
    resp = google.get("/oauth2/v2/userinfo")
    if not resp.ok:
        return "Google auth failed", 401
    info = resp.json()
    email = info["email"]
    username = info.get("name", "GoogleUser")

    users = db_ref.child("users").get() or {}
    for key, user in users.items():
        if user["email"] == email:
            session.update({
                'user_id': key,
                'user': email,
                'username': user.get("username"),
                'role': user.get("role", "user"),
                'department': user.get("department", "NotSet")
            })
            break
    else:
        new_user = db_ref.child("users").push({
            "username": username,
            "email": email,
            "password": "",
            "role": "user",
            "department": "NotSet"
        })
        session.update({
            'user_id': new_user.key,
            'user': email,
            'username': username,
            'role': "user",
            'department': "NotSet"
        })
    if session["department"] == "NotSet":
        return redirect("/select-department")
    return redirect("/dashboard")

@app.route('/dashboard')
def dashboard():
    user_id = session.get("user_id")
    if not user_id:
        return redirect("/login")
    user_data = db_ref.child("users").child(user_id).get()
    if not user_data or user_data.get("department") == "NotSet":
        return redirect("/select-department")
    return render_template("dashboard.html", name=user_data["username"], email=user_data["email"], department=user_data["department"])

@app.route("/select-department", methods=["GET", "POST"])
def select_department():
    if request.method == "POST":
        department = request.form.get("department")
        user_id = session.get("user_id")
        if user_id and department:
            db_ref.child("users").child(user_id).update({"department": department})
            session["department"] = department
            return redirect("/dashboard")
    return render_template("select_department.html")

@app.route('/admin', methods=['GET', 'POST'])
def admin_panel():
    if session.get('role') != 'admin':
        return redirect('/login')
    users = db_ref.child("users").get() or {}
    user_list = [
        {"key": key, **val}
        for key, val in users.items()
    ]
    task_data = db_ref.child("department_tasks").get() or {}
    selected_dept = clean_department(request.form.get("department")) if request.method == 'POST' else None
    task_text = "\n".join(task_data.get(selected_dept, [])) if selected_dept else ""
    return render_template("admin.html", users=user_list, task_data=task_data, selected_dept=selected_dept, task_text=task_text)

@app.route('/update_user/<user_key>', methods=['POST'])
def update_user(user_key):
    if session.get('role') != 'admin':
        return redirect('/login')
    new_dept = clean_department(request.form.get('department'))
    db_ref.child("users").child(user_key).update({"department": new_dept})
    return redirect("/admin")

@app.route('/delete_user/<user_key>')
def delete_user(user_key):
    if session.get('role') == 'admin':
        db_ref.child("users").child(user_key).delete()
    return redirect("/admin")

@app.route('/update_tasks', methods=['POST'])
def update_tasks():
    department = clean_department(request.form.get('department'))
    tasks_text = request.form.get('tasks')
    tasks = [t.strip() for t in tasks_text.splitlines() if t.strip()]
    db_ref.child("department_tasks").child(department).set(tasks)
    return redirect('/admin')

@app.route("/calculate-time", methods=["GET", "POST"])
def calculate_time():
    result = None
    if request.method == "POST":
        try:
            load_cycles = int(request.form["load_cycles"])
            environment = request.form["environment"]
            speed = float(request.form["speed"])
            input_df = pd.DataFrame([{
                "Engine Hours": 1500,
                "Fuel Used (L)": 5.0,
                "Load Cycles": load_cycles,
                "Idling Time (min)": 30,
                "Proximity Hazard Count": 1,
                "speed": speed,
                "Machine ID": "EXC001",
                "Operator ID": "OP1001",
                "Environmental Conditions": environment,
                "Maintenance Flag": "No"
            }])
            time_sec = model.predict(input_df)[0]
            adjusted = time_sec * adjustment_factors.get(environment, 1.0)
            result = round(adjusted / 60, 2)
        except Exception:
            result = "Error"
    return render_template("time.html", result=result)

@app.route('/manual', methods=['POST'])
def manual():
    try:
        df = pd.DataFrame([{
            'Load Cycles': int(request.form['load_cycles']),
            'Idling Time (min)': float(request.form['idling_time']),
            'Fuel Used (L)': float(request.form['fuel_used']),
            'Environmental Conditions': request.form['environmental_conditions'],
            'Safety Alert Triggered': request.form['safety_alert']
        }])
        df['Estimated_Task_Time'] = (df['Load Cycles'] * 5) + df['Idling Time (min)'] + (df['Fuel Used (L)'] * 10)
        df['Adjusted_Task_Time'] = df.apply(lambda row: row['Estimated_Task_Time'] * (1.2 if row['Environmental Conditions'] in ['Foggy', 'Snowy', 'Rainy'] else 1.0) * (1.1 if row['Safety Alert Triggered'].strip().lower() == 'yes' else 1.0), axis=1)
        row = df.iloc[0]
        result = f"<b>Estimated:</b> {row['Estimated_Task_Time']:.2f} min<br><b>Adjusted:</b> {row['Adjusted_Task_Time']:.2f} min"
        return render_template("time.html", result=result)
    except Exception as e:
        return str(e), 400

@app.route('/upload', methods=['POST'])
def upload():
    file = request.files.get('file')
    if not file or not allowed_file(file.filename):
        return "Invalid", 400
    filename = secure_filename(file.filename)
    file.save(os.path.join(app.config['UPLOAD_FOLDER'], filename))
    return redirect('/dashboard')

@app.route('/summarize/<filename>', methods=['POST'])
def summarize_document(filename):
    path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    summary = summarize_text(extract_text(path))
    summary_path = os.path.join(app.config['SUMMARY_FOLDER'], f"{filename}_summary.txt")
    with open(summary_path, 'w', encoding='utf-8') as f:
        f.write(summary)
    return jsonify({'summary': summary, 'download': f"/download/{filename}_summary.txt"})

@app.route('/ask/<filename>', methods=['POST'])
def ask_question(filename):
    try:
        question = request.json.get("question", "")
        content = extract_text(os.path.join(app.config['UPLOAD_FOLDER'], filename))[:12000]
        prompt = f"Based on this document:\n{content}\n\nQuestion: {question}\nAnswer:"
        response = genai.GenerativeModel("models/gemini-1.5-flash").generate_content(prompt)
        return jsonify({"answer": response.text.strip()})
    except Exception as e:
        return jsonify({"answer": "⚠️ Error"}), 500

@app.route('/download/<filename>')
def download(filename):
    return send_from_directory(app.config['SUMMARY_FOLDER'], filename, as_attachment=True)

@app.route('/delete/<filename>', methods=['POST'])
def delete_file(filename):
    try:
        os.remove(os.path.join(app.config['UPLOAD_FOLDER'], filename))
        os.remove(os.path.join(app.config['SUMMARY_FOLDER'], f"{filename}_summary.txt"))
        return jsonify({'status': 'success'})
    except:
        return jsonify({'status': 'error'}), 500

@app.route('/ask-bot')
def ask_bot():
    files = os.listdir(app.config['UPLOAD_FOLDER'])
    return render_template('doubt-bot.html', files=files)

@app.route('/video-summary', methods=['GET', 'POST'])
def video_summary():
    category = request.form.get("category", "wheel_tractor")
    action = request.form.get("action", "view")
    summary = summarize(category) if action == "summarize" else None
    return render_template_string(TEMPLATE, videos=video_data[category], category=category, summary=summary)

@app.route('/contact')
def contact():
    return render_template('emailjs.html')

@app.route('/safety')
def safety():
    return render_template('safety.html')

@app.route('/log', methods=['POST'])
def log_issue():
    data = request.get_json()
    log = f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] Operator: {data.get('operator')}, Machine ID: {data.get('machineId')}, Issue: {data.get('issue')}\n"
    with open('logs.txt', 'a') as f:
        f.write(log)
    return "Log saved", 200

@app.route('/get_logs')
def get_logs_by_machine():
    machine_id = request.args.get('machineId', '').strip()
    logs = []
    try:
        with open('logs.txt', 'r') as f:
            logs = [line.strip() for line in f if f"Machine ID: {machine_id}" in line]
    except FileNotFoundError:
        pass
    return jsonify({'logs': logs})

if __name__ == '__main__':
    app.run(debug=True)
